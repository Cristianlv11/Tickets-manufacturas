# -*- coding: utf-8 -*-
"""
Módulo de Gestión TI: copias de seguridad, hoja de vida de activos,
bóveda de contraseñas y vencimiento de dominios.
"""
import csv
import io
import json
import os
import re
import secrets
from datetime import datetime, date, timedelta, timezone
from typing import Optional

import openpyxl
from cryptography.fernet import Fernet, InvalidToken
from fastapi import APIRouter, Depends, Form, HTTPException, Request, Response, UploadFile, File, status
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse
from fastapi.templating import Jinja2Templates
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from sqlalchemy.orm import Session

import models
from database import get_db
from equipos_seed import EQUIPOS_OLAMTEX
from security import (
    COOKIE_SECURE, DEFAULT_ADMIN_PASSWORD, create_access_token, decode_token,
    get_current_admin, get_password_hash, verify_password,
)

router = APIRouter()
templates = Jinja2Templates(directory="templates")

COLOMBIA = timezone(timedelta(hours=-5))
TIPOS_ESTACION = ("Puesto de Trabajo", "Laptop", "Servidor")
TIPOS_RESPALDO = ("Físico", "Nube")
CATEGORIAS_CREDENCIAL = ["Correo", "Software", "Aplicativo", "Cámara / CCTV", "Hosting / Dominio", "Red / Wi-Fi", "Equipo", "Otro"]
ESTADOS_ACTIVO = ["Activo", "En reparación", "En bodega", "De baja"]
CATEGORIAS_ACTIVO = ["Mobiliario", "Maquinaria", "Equipo de oficina", "Herramienta", "Electrodoméstico", "Tecnología", "Vehículo", "Otro"]
FOTO_MAX_BYTES = 10 * 1024 * 1024
FOTO_MAX_POR_SUBIDA = 10
BOVEDA_COOKIE = "boveda_unlock"
BOVEDA_PATH = "/admin/boveda"
BOVEDA_MINUTOS = 10


def ahora():
    """Hora de Colombia (UTC-5, sin horario de verano) como datetime sin zona."""
    return datetime.now(COLOMBIA).replace(tzinfo=None)


def hoy():
    return ahora().date()


def redirect(url):
    return RedirectResponse(url=url, status_code=status.HTTP_302_FOUND)


def parse_fecha(valor: Optional[str]):
    if not valor:
        return None
    try:
        return date.fromisoformat(valor[:10])
    except ValueError:
        return None


def parse_int(valor):
    try:
        return int(str(valor).replace(".", "").replace(",", "").replace("$", "").strip())
    except (TypeError, ValueError):
        return None


def cargar_equipos(db: Session):
    config = db.query(models.RedConfig).filter(models.RedConfig.clave == "dispositivos").first()
    if config and config.valor:
        try:
            return json.loads(config.valor)
        except Exception:
            pass
    return EQUIPOS_OLAMTEX


def general_a_equipo(activo: models.ActivoGeneral):
    """Presenta un activo general con las mismas claves que un equipo del inventario de red."""
    return {
        "id": f"ag-{activo.id}", "general": True, "nombre": activo.nombre or "",
        "tipo": activo.categoria or "Activo general", "piso": activo.ubicacion or "",
        "usuario": activo.responsable or "", "cargo": "", "notas": activo.notas or "",
    }


def id_general(equipo_id: str):
    """'ag-12' -> 12; None si no es un activo general."""
    if equipo_id.startswith("ag-") and equipo_id[3:].isdigit():
        return int(equipo_id[3:])
    return None


def cargar_activos(db: Session):
    """Equipos del inventario de red más los activos generales."""
    generales = db.query(models.ActivoGeneral).order_by(models.ActivoGeneral.id).all()
    return cargar_equipos(db) + [general_a_equipo(a) for a in generales]


def guardar_fotos(db: Session, equipo_id: str, archivos, descripcion: str, usuario: str):
    """Guarda las imágenes subidas; ignora archivos vacíos, que no sean imagen o demasiado grandes."""
    guardadas = 0
    for f in (archivos or [])[:FOTO_MAX_POR_SUBIDA]:
        if not f or not f.filename:
            continue
        contenido = f.file.read()
        tipo = (f.content_type or "").lower()
        if not contenido or len(contenido) > FOTO_MAX_BYTES or not tipo.startswith("image/"):
            continue
        db.add(models.ActivoFoto(
            equipo_id=equipo_id, nombre_archivo=f.filename, tipo_contenido=tipo, archivo_bytes=contenido,
            tamano=len(contenido), descripcion=(descripcion or "").strip() or None,
            subido_por=usuario, fecha_subida=ahora(),
        ))
        guardadas += 1
    return guardadas


def es_estacion(equipo):
    return any(t in (equipo.get("tipo") or "") for t in TIPOS_ESTACION)


def tamano_legible(n):
    n = n or 0
    for unidad in ("B", "KB", "MB", "GB", "TB"):
        if n < 1024:
            return f"{n:.0f} {unidad}" if unidad == "B" else f"{n:.1f} {unidad}"
        n /= 1024
    return f"{n:.1f} PB"


def hace_cuanto(fecha):
    if not fecha:
        return "Nunca"
    seg = int((ahora() - fecha).total_seconds())
    if seg < 3600:
        return f"hace {max(seg // 60, 0)} min"
    if seg < 86400:
        return f"hace {seg // 3600} h"
    return f"hace {seg // 86400} días"


# ---------------------------------------------------------------------------
# Copias de seguridad
# ---------------------------------------------------------------------------

NIVEL_ORDEN = {"ok": 0, "no_aplica": 0, "atraso": 1, "sin_datos": 2, "vencido": 3, "error": 3}


def estado_respaldo(estacion: models.RespaldoEstacion, tipo: str):
    requerido = estacion.requiere_fisico if tipo == "Físico" else estacion.requiere_nube
    registros = [r for r in estacion.registros if r.tipo == tipo]
    ultimo = registros[0] if registros else None
    ultimo_ok = next((r for r in registros if r.estado == "OK"), None)

    if not requerido and not ultimo:
        nivel = "no_aplica"
    elif not ultimo_ok:
        nivel = "error" if ultimo else "sin_datos"
    elif ultimo.estado != "OK":
        nivel = "error"
    else:
        horas = (ahora() - ultimo_ok.fecha).total_seconds() / 3600
        frecuencia = estacion.frecuencia_horas or 24
        if horas <= frecuencia + 6:
            nivel = "ok"
        elif horas <= frecuencia * 2 + 6:
            nivel = "atraso"
        else:
            nivel = "vencido"
    if not requerido and nivel in ("sin_datos", "vencido", "atraso"):
        nivel = "no_aplica"
    return {"nivel": nivel, "ultimo": ultimo, "ultimo_ok": ultimo_ok, "requerido": requerido}


def resumen_estacion(estacion):
    fisico = estado_respaldo(estacion, "Físico")
    nube = estado_respaldo(estacion, "Nube")
    nivel = max((fisico["nivel"], nube["nivel"]), key=lambda n: NIVEL_ORDEN[n])
    if not estacion.activo:
        nivel = "no_aplica"
    return {"estacion": estacion, "fisico": fisico, "nube": nube, "nivel": nivel}


def obtener_token_respaldos(db: Session, regenerar=False):
    config = db.query(models.RedConfig).filter(models.RedConfig.clave == "respaldos_api_token").first()
    if not config or regenerar:
        token = secrets.token_urlsafe(32)
        if not config:
            config = models.RedConfig(clave="respaldos_api_token", valor=token)
            db.add(config)
        else:
            config.valor = token
        db.commit()
    return config.valor


@router.get("/admin/respaldos", response_class=HTMLResponse)
def respaldos_page(request: Request, user=Depends(get_current_admin), db: Session = Depends(get_db)):
    if not user:
        return redirect("/login")
    estaciones = db.query(models.RespaldoEstacion).order_by(models.RespaldoEstacion.nombre).all()
    resumen = [resumen_estacion(e) for e in estaciones]
    resumen.sort(key=lambda r: (-NIVEL_ORDEN[r["nivel"]], r["estacion"].nombre or ""))
    conteo = {
        "ok": sum(1 for r in resumen if r["nivel"] == "ok"),
        "atraso": sum(1 for r in resumen if r["nivel"] == "atraso"),
        "problema": sum(1 for r in resumen if r["nivel"] in ("vencido", "error", "sin_datos")),
        "total": sum(1 for r in resumen if r["estacion"].activo),
    }
    return templates.TemplateResponse(request=request, name="ti_respaldos.html", context={
        "request": request, "user": user, "resumen": resumen, "conteo": conteo,
        "token": obtener_token_respaldos(db), "tamano_legible": tamano_legible, "hace_cuanto": hace_cuanto,
        "ahora": ahora(), "base_url": str(request.base_url).rstrip("/"),
    })


@router.post("/admin/respaldos/sincronizar")
def respaldos_sincronizar(user=Depends(get_current_admin), db: Session = Depends(get_db)):
    if not user:
        return redirect("/login")
    existentes = {e.hostname for e in db.query(models.RespaldoEstacion).all()}
    for eq in cargar_equipos(db):
        hostname = (eq.get("hostname") or "").strip().upper()
        if es_estacion(eq) and hostname and hostname not in existentes:
            db.add(models.RespaldoEstacion(
                hostname=hostname, equipo_id=eq.get("id"),
                nombre=f'{eq.get("nombre", hostname)} — {eq.get("usuario", "")}'.strip(" —"),
                frecuencia_horas=24, requiere_fisico=True, requiere_nube=False, activo=True,
            ))
            existentes.add(hostname)
    db.commit()
    return redirect("/admin/respaldos")


@router.post("/admin/respaldos/estacion")
def respaldos_guardar_estacion(
    estacion_id: Optional[int] = Form(None),
    hostname: str = Form(...),
    nombre: str = Form(""),
    frecuencia_horas: int = Form(24),
    requiere_fisico: Optional[str] = Form(None),
    requiere_nube: Optional[str] = Form(None),
    activo: Optional[str] = Form(None),
    notas: str = Form(""),
    user=Depends(get_current_admin), db: Session = Depends(get_db),
):
    if not user:
        return redirect("/login")
    hostname = hostname.strip().upper()
    estacion = db.get(models.RespaldoEstacion, estacion_id) if estacion_id else None
    if not estacion:
        estacion = db.query(models.RespaldoEstacion).filter(models.RespaldoEstacion.hostname == hostname).first()
    if not estacion:
        estacion = models.RespaldoEstacion(hostname=hostname)
        db.add(estacion)
    estacion.hostname = hostname
    estacion.nombre = nombre.strip() or hostname
    estacion.frecuencia_horas = max(1, frecuencia_horas)
    estacion.requiere_fisico = bool(requiere_fisico)
    estacion.requiere_nube = bool(requiere_nube)
    estacion.activo = bool(activo)
    estacion.notas = notas.strip() or None
    eq = next((e for e in cargar_equipos(db) if (e.get("hostname") or "").upper() == hostname), None)
    if eq:
        estacion.equipo_id = eq.get("id")
    db.commit()
    return redirect("/admin/respaldos")


@router.post("/admin/respaldos/estacion/{estacion_id}/delete")
def respaldos_eliminar_estacion(estacion_id: int, user=Depends(get_current_admin), db: Session = Depends(get_db)):
    if not user:
        return redirect("/login")
    estacion = db.get(models.RespaldoEstacion, estacion_id)
    if estacion:
        db.delete(estacion)
        db.commit()
    return redirect("/admin/respaldos")


@router.post("/admin/respaldos/manual")
def respaldos_registro_manual(
    estacion_id: int = Form(...),
    tipo: str = Form(...),
    estado: str = Form("OK"),
    fecha: str = Form(""),
    destino: str = Form(""),
    mensaje: str = Form(""),
    user=Depends(get_current_admin), db: Session = Depends(get_db),
):
    if not user:
        return redirect("/login")
    estacion = db.get(models.RespaldoEstacion, estacion_id)
    if estacion and tipo in TIPOS_RESPALDO:
        try:
            fecha_dt = datetime.fromisoformat(fecha) if fecha else ahora()
        except ValueError:
            fecha_dt = ahora()
        db.add(models.RespaldoRegistro(
            estacion_id=estacion.id, tipo=tipo, estado="OK" if estado == "OK" else "Error",
            fecha=fecha_dt, destino=destino.strip() or None, mensaje=mensaje.strip() or None,
            origen="Manual", registrado_por=user.username,
        ))
        db.commit()
    return redirect("/admin/respaldos")


@router.post("/admin/respaldos/token/regenerar")
def respaldos_regenerar_token(user=Depends(get_current_admin), db: Session = Depends(get_db)):
    if not user:
        return redirect("/login")
    obtener_token_respaldos(db, regenerar=True)
    return redirect("/admin/respaldos")


@router.get("/admin/respaldos/script")
def respaldos_descargar_script(request: Request, user=Depends(get_current_admin), db: Session = Depends(get_db)):
    if not user:
        return redirect("/login")
    with open(os.path.join("scripts", "respaldo_olamtex.ps1"), encoding="utf-8") as f:
        contenido = f.read()
    base_url = str(request.base_url).rstrip("/")
    if request.headers.get("x-forwarded-proto") == "https":
        base_url = base_url.replace("http://", "https://", 1)
    contenido = contenido.replace("__URL_APP__", base_url).replace("__TOKEN__", obtener_token_respaldos(db))
    # BOM UTF-8 para que Windows PowerShell 5.1 lea bien las tildes
    return Response(
        content="﻿" + contenido.replace("\r\n", "\n").replace("\n", "\r\n"),
        media_type="text/plain; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="respaldo_olamtex.ps1"'},
    )


@router.post("/api/respaldos/reporte")
async def respaldos_api_reporte(request: Request, db: Session = Depends(get_db)):
    token = request.headers.get("x-backup-token", "")
    if not secrets.compare_digest(token, obtener_token_respaldos(db)):
        raise HTTPException(status_code=401, detail="Token inválido")
    crudo = await request.body()
    try:
        texto = crudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = crudo.decode("cp1252", errors="replace")  # PowerShell antiguo puede enviar ANSI
    try:
        body = json.loads(texto)
    except ValueError:
        raise HTTPException(status_code=400, detail="JSON inválido")
    if not isinstance(body, dict):
        raise HTTPException(status_code=400, detail="JSON inválido")

    hostname = str(body.get("hostname") or "").strip().upper()
    tipo = "Nube" if str(body.get("tipo", "")).lower().startswith("nube") else "Físico"
    if not hostname:
        raise HTTPException(status_code=400, detail="Falta hostname")

    estacion = db.query(models.RespaldoEstacion).filter(models.RespaldoEstacion.hostname == hostname).first()
    if not estacion:
        eq = next((e for e in cargar_equipos(db) if (e.get("hostname") or "").upper() == hostname), None)
        estacion = models.RespaldoEstacion(
            hostname=hostname, equipo_id=eq.get("id") if eq else None,
            nombre=(f'{eq.get("nombre")} — {eq.get("usuario", "")}' if eq else hostname),
            frecuencia_horas=24, requiere_fisico=(tipo == "Físico"), requiere_nube=(tipo == "Nube"), activo=True,
        )
        db.add(estacion)
        db.flush()

    fecha = ahora()
    if body.get("fecha"):
        try:
            f = datetime.fromisoformat(str(body["fecha"]))
            fecha = f.astimezone(COLOMBIA).replace(tzinfo=None) if f.tzinfo else f
        except ValueError:
            pass

    registro = models.RespaldoRegistro(
        estacion_id=estacion.id, tipo=tipo,
        estado="OK" if str(body.get("estado", "")).upper() == "OK" else "Error",
        fecha=fecha, destino=str(body.get("destino") or "")[:500] or None,
        tamano_bytes=parse_int(body.get("tamano_bytes")) or 0,
        archivos=parse_int(body.get("archivos")) or 0,
        duracion_seg=parse_int(body.get("duracion_seg")) or 0,
        mensaje=str(body.get("mensaje") or "")[:4000] or None,
        origen="Script",
    )
    db.add(registro)
    db.commit()
    return {"status": "ok", "estacion": estacion.hostname, "registro_id": registro.id}


@router.get("/admin/rustdesk/script")
def rustdesk_descargar_script(request: Request, user=Depends(get_current_admin), db: Session = Depends(get_db)):
    if not user:
        return redirect("/login")
    with open(os.path.join("scripts", "instalar_rustdesk.ps1"), encoding="utf-8") as f:
        contenido = f.read()
    base_url = str(request.base_url).rstrip("/")
    if request.headers.get("x-forwarded-proto") == "https":
        base_url = base_url.replace("http://", "https://", 1)
    contenido = contenido.replace("__URL_APP__", base_url).replace("__TOKEN__", obtener_token_respaldos(db))
    return Response(
        content="﻿" + contenido.replace("\r\n", "\n").replace("\n", "\r\n"),
        media_type="text/plain; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="instalar_rustdesk.ps1"'},
    )


@router.post("/api/boveda/rustdesk")
async def rustdesk_api_registrar(request: Request, db: Session = Depends(get_db)):
    """El instalador de cada PC deja aquí su ID y contraseña de RustDesk (solo crea/actualiza, nunca lee)."""
    token = request.headers.get("x-backup-token", "")
    if not secrets.compare_digest(token, obtener_token_respaldos(db)):
        raise HTTPException(status_code=401, detail="Token inválido")
    if not fernet():
        raise HTTPException(status_code=503, detail="Bóveda no configurada")
    try:
        body = json.loads((await request.body()).decode("utf-8-sig"))
    except (ValueError, UnicodeDecodeError):
        raise HTTPException(status_code=400, detail="JSON inválido")
    hostname = str(body.get("hostname") or "").strip().upper()[:60]
    rid = re.sub(r"\D", "", str(body.get("id") or ""))
    clave = str(body.get("password") or "")
    if not hostname or not rid or len(clave) < 8:
        raise HTTPException(status_code=400, detail="Faltan hostname, id o password")

    eq = next((e for e in cargar_equipos(db) if (e.get("hostname") or "").upper() == hostname), None)
    servicio = f"RustDesk · {hostname}" + (f" ({eq.get('usuario')})" if eq and eq.get("usuario") else "")
    credencial = db.query(models.Credencial).filter(models.Credencial.servicio.like(f"RustDesk · {hostname}%")).first()
    nueva = credencial is None
    if nueva:
        credencial = models.Credencial(fecha_creacion=ahora())
        db.add(credencial)
    credencial.categoria = "Equipo"
    credencial.servicio = servicio
    credencial.usuario = rid
    credencial.clave_cifrada = cifrar(clave)
    credencial.equipo_id = eq.get("id") if eq else None
    credencial.notas_cifradas = cifrar(
        f"ID RustDesk: {rid} | Instalado/actualizado por script el {ahora().strftime('%d/%m/%Y %I:%M %p')}"
        f" | Versión: {str(body.get('version') or '')[:20]} | Usuario de Windows: {str(body.get('usuario_windows') or '')[:60]}")
    credencial.actualizado_por = "script RustDesk"
    credencial.fecha_actualizacion = ahora()
    db.flush()
    db.add(models.CredencialAcceso(credencial_id=credencial.id, servicio=servicio, usuario_app="script RustDesk",
                                   accion="Crear" if nueva else "Editar",
                                   ip=(request.headers.get("x-forwarded-for") or "")[:100], fecha=ahora()))
    db.commit()
    return {"status": "ok", "servicio": servicio, "nueva": nueva}


# ---------------------------------------------------------------------------
# Control de activos (hoja de vida por equipo)
# ---------------------------------------------------------------------------

def consolidar_activos(db: Session):
    equipos = cargar_activos(db)
    fichas = {f.equipo_id: f for f in db.query(models.ActivoFicha).all()}
    fotos = {}
    for foto_id, eid in db.query(models.ActivoFoto.id, models.ActivoFoto.equipo_id).order_by(models.ActivoFoto.id).all():
        fotos.setdefault(eid, []).append(foto_id)
    mantenimientos = {}
    for m in db.query(models.ActivoMantenimiento).order_by(models.ActivoMantenimiento.fecha.desc()).all():
        mantenimientos.setdefault(m.equipo_id, []).append(m)
    software = {}
    for s in db.query(models.ActivoSoftware).all():
        software.setdefault(s.equipo_id, []).append(s)
    respaldos = {e.equipo_id: resumen_estacion(e) for e in db.query(models.RespaldoEstacion).all() if e.equipo_id}

    filas = []
    for eq in equipos:
        eid = eq.get("id")
        ficha = fichas.get(eid)
        mants = mantenimientos.get(eid, [])
        ultimo_mant = mants[0] if mants else None
        resp = respaldos.get(eid)
        garantia = ficha.garantia_hasta if ficha else None
        filas.append({
            "id": eid,
            "nombre": eq.get("nombre", ""),
            "tipo": eq.get("tipo", ""),
            "piso": eq.get("piso", ""),
            "usuario": eq.get("usuario", ""),
            "cargo": eq.get("cargo", ""),
            "hostname": eq.get("hostname", ""),
            "ip": eq.get("ip", ""),
            "marca": eq.get("marca", ""),
            "so": eq.get("so", ""),
            "hardware": eq.get("hardware", ""),
            "serial": ficha.serial if ficha else "",
            "placa": ficha.placa_inventario if ficha else "",
            "estado": ficha.estado if ficha else "Activo",
            "garantia_hasta": garantia.isoformat() if garantia else "",
            "garantia_vigente": bool(garantia and garantia >= hoy()),
            "ultimo_mantenimiento": ultimo_mant.fecha.isoformat() if ultimo_mant and ultimo_mant.fecha else "",
            "ultimo_mantenimiento_tipo": ultimo_mant.tipo if ultimo_mant else "",
            "dias_sin_mantenimiento": (hoy() - ultimo_mant.fecha).days if ultimo_mant and ultimo_mant.fecha else None,
            "total_mantenimientos": len(mants),
            "total_software": len({s.nombre for s in software.get(eid, []) if s.accion != "Desinstalado"}),
            "respaldo": resp["nivel"] if resp else "",
            "general": bool(eq.get("general")),
            "fotos": len(fotos.get(eid, [])),
            "foto_id": fotos[eid][0] if fotos.get(eid) else None,
        })
    return filas


@router.get("/admin/activos", response_class=HTMLResponse)
def activos_page(request: Request, user=Depends(get_current_admin), db: Session = Depends(get_db)):
    if not user:
        return redirect("/login")
    filas = consolidar_activos(db)
    return templates.TemplateResponse(request=request, name="ti_activos.html", context={
        "request": request, "user": user, "filas": filas,
        "pisos": sorted({f["piso"] for f in filas if f["piso"]}),
        "tipos": sorted({f["tipo"] for f in filas if f["tipo"]}),
        "estados": ESTADOS_ACTIVO, "categorias": CATEGORIAS_ACTIVO,
    })


@router.post("/admin/activos/nuevo")
def activo_crear_general(
    request: Request,
    nombre: str = Form(...), categoria: str = Form(""), ubicacion: str = Form(""), responsable: str = Form(""),
    placa_inventario: str = Form(""), notas: str = Form(""), fotos: list[UploadFile] = File(None),
    user=Depends(get_current_admin), db: Session = Depends(get_db),
):
    if not user:
        return redirect("/login")
    activo = models.ActivoGeneral(
        nombre=nombre.strip(), categoria=categoria.strip() or None, ubicacion=ubicacion.strip() or None,
        responsable=responsable.strip() or None, notas=notas.strip() or None,
        creado_por=user.username, fecha_creacion=ahora(),
    )
    db.add(activo)
    db.flush()
    equipo_id = f"ag-{activo.id}"
    db.add(models.ActivoFicha(equipo_id=equipo_id, placa_inventario=placa_inventario.strip() or None, estado="Activo"))
    guardar_fotos(db, equipo_id, fotos, "", user.username)
    db.commit()
    if "application/json" in (request.headers.get("accept") or ""):
        return {"status": "ok", "id": equipo_id}
    return redirect(f"/admin/activos/{equipo_id}")


@router.get("/admin/activos/foto/{foto_id}")
def activo_ver_foto(foto_id: int, user=Depends(get_current_admin), db: Session = Depends(get_db)):
    if not user:
        return redirect("/login")
    foto = db.get(models.ActivoFoto, foto_id)
    if not foto:
        raise HTTPException(status_code=404)
    return Response(content=foto.archivo_bytes, media_type=foto.tipo_contenido or "image/jpeg",
                    headers={"Cache-Control": "private, max-age=86400"})


@router.get("/admin/activos/exportar")
def activos_exportar(user=Depends(get_current_admin), db: Session = Depends(get_db)):
    if not user:
        return redirect("/login")
    filas = consolidar_activos(db)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Activos"
    columnas = [
        ("Equipo", "nombre"), ("Tipo", "tipo"), ("Piso", "piso"), ("Usuario", "usuario"), ("Cargo", "cargo"),
        ("Hostname", "hostname"), ("IP", "ip"), ("Marca", "marca"), ("Sistema operativo", "so"),
        ("Hardware", "hardware"), ("Serial", "serial"), ("Placa inventario", "placa"), ("Estado", "estado"),
        ("Garantía hasta", "garantia_hasta"), ("Último mantenimiento", "ultimo_mantenimiento"),
        ("Tipo último mant.", "ultimo_mantenimiento_tipo"), ("N.º mantenimientos", "total_mantenimientos"),
        ("Programas registrados", "total_software"), ("Estado respaldo", "respaldo"), ("Fotos", "fotos"),
    ]
    ws.append([c[0] for c in columnas])
    _estilo_encabezado(ws)
    for f in filas:
        ws.append([f[c[1]] for c in columnas])
    _autoajustar(ws)

    ws2 = wb.create_sheet("Mantenimientos")
    ws2.append(["Equipo", "Fecha", "Tipo", "Descripción", "Técnico", "Ticket", "Costo"])
    _estilo_encabezado(ws2)
    nombres = {f["id"]: f["nombre"] for f in filas}
    for m in db.query(models.ActivoMantenimiento).order_by(models.ActivoMantenimiento.fecha.desc()).all():
        ws2.append([nombres.get(m.equipo_id, m.equipo_id), m.fecha, m.tipo, m.descripcion, m.tecnico, m.ticket_id, m.costo])
    _autoajustar(ws2)

    ws3 = wb.create_sheet("Software")
    ws3.append(["Equipo", "Fecha", "Acción", "Programa", "Versión", "Licencia", "Notas"])
    _estilo_encabezado(ws3)
    for s in db.query(models.ActivoSoftware).order_by(models.ActivoSoftware.fecha.desc()).all():
        ws3.append([nombres.get(s.equipo_id, s.equipo_id), s.fecha, s.accion, s.nombre, s.version, s.licencia, s.notas])
    _autoajustar(ws3)

    stream = io.BytesIO()
    wb.save(stream)
    nombre = f"Activos TI Olamtex {ahora().strftime('%Y-%m-%d')}.xlsx"
    return Response(
        content=stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre}"'},
    )


def _estilo_encabezado(ws):
    borde = Side(style="thin", color="E2E8F0")
    for cell in ws[1]:
        cell.fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
        cell.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=borde, right=borde, top=borde, bottom=borde)
    ws.freeze_panes = "A2"


def _autoajustar(ws):
    for col in ws.columns:
        largo = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = max(min(largo + 3, 50), 12)
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


@router.get("/admin/activos/{equipo_id}", response_class=HTMLResponse)
def activo_detalle(request: Request, equipo_id: str, user=Depends(get_current_admin), db: Session = Depends(get_db)):
    if not user:
        return redirect("/login")
    equipo = next((e for e in cargar_activos(db) if e.get("id") == equipo_id), None)
    if not equipo:
        raise HTTPException(status_code=404, detail="Equipo no encontrado")
    ficha = db.query(models.ActivoFicha).filter(models.ActivoFicha.equipo_id == equipo_id).first()
    fotos = db.query(models.ActivoFoto.id, models.ActivoFoto.descripcion, models.ActivoFoto.fecha_subida,
                     models.ActivoFoto.subido_por, models.ActivoFoto.tamano) \
        .filter(models.ActivoFoto.equipo_id == equipo_id).order_by(models.ActivoFoto.id).all()
    mantenimientos = db.query(models.ActivoMantenimiento).filter(models.ActivoMantenimiento.equipo_id == equipo_id) \
        .order_by(models.ActivoMantenimiento.fecha.desc(), models.ActivoMantenimiento.id.desc()).all()
    software = db.query(models.ActivoSoftware).filter(models.ActivoSoftware.equipo_id == equipo_id) \
        .order_by(models.ActivoSoftware.fecha.desc(), models.ActivoSoftware.id.desc()).all()
    estacion = db.query(models.RespaldoEstacion).filter(models.RespaldoEstacion.equipo_id == equipo_id).first()
    return templates.TemplateResponse(request=request, name="ti_activo_detalle.html", context={
        "request": request, "user": user, "equipo": equipo, "ficha": ficha, "fotos": fotos,
        "mantenimientos": mantenimientos, "software": software,
        "respaldo": resumen_estacion(estacion) if estacion else None,
        "estados": ESTADOS_ACTIVO, "categorias": CATEGORIAS_ACTIVO,
        "hoy": hoy(), "hace_cuanto": hace_cuanto, "tamano_legible": tamano_legible,
    })


@router.post("/admin/activos/{equipo_id}/fotos")
def activo_subir_fotos(
    equipo_id: str, request: Request, fotos: list[UploadFile] = File(...), descripcion: str = Form(""),
    user=Depends(get_current_admin), db: Session = Depends(get_db),
):
    if not user:
        return redirect("/login")
    if not any(e.get("id") == equipo_id for e in cargar_activos(db)):
        raise HTTPException(status_code=404, detail="Equipo no encontrado")
    guardadas = guardar_fotos(db, equipo_id, fotos, descripcion, user.username)
    db.commit()
    if "application/json" in (request.headers.get("accept") or ""):
        return {"status": "ok", "guardadas": guardadas}
    return redirect(f"/admin/activos/{equipo_id}#fotos")


@router.post("/admin/activos/{equipo_id}/general")
def activo_editar_general(
    equipo_id: str,
    nombre: str = Form(...), categoria: str = Form(""), ubicacion: str = Form(""), responsable: str = Form(""),
    notas: str = Form(""), user=Depends(get_current_admin), db: Session = Depends(get_db),
):
    if not user:
        return redirect("/login")
    activo = db.get(models.ActivoGeneral, id_general(equipo_id) or 0)
    if not activo:
        raise HTTPException(status_code=404)
    activo.nombre = nombre.strip()
    activo.categoria = categoria.strip() or None
    activo.ubicacion = ubicacion.strip() or None
    activo.responsable = responsable.strip() or None
    activo.notas = notas.strip() or None
    db.commit()
    return redirect(f"/admin/activos/{equipo_id}")


@router.post("/admin/activos/{equipo_id}/eliminar")
def activo_eliminar_general(equipo_id: str, user=Depends(get_current_admin), db: Session = Depends(get_db)):
    """Solo para activos generales; los equipos de red se eliminan en Administración de Equipos."""
    if not user:
        return redirect("/login")
    activo = db.get(models.ActivoGeneral, id_general(equipo_id) or 0)
    if not activo:
        raise HTTPException(status_code=404)
    for modelo in (models.ActivoFicha, models.ActivoFoto, models.ActivoMantenimiento, models.ActivoSoftware):
        db.query(modelo).filter(modelo.equipo_id == equipo_id).delete(synchronize_session=False)
    db.delete(activo)
    db.commit()
    return redirect("/admin/activos")


@router.post("/admin/activos/{equipo_id}/ficha")
def activo_guardar_ficha(
    equipo_id: str,
    placa_inventario: str = Form(""), serial: str = Form(""), estado: str = Form("Activo"),
    fecha_compra: str = Form(""), proveedor: str = Form(""), garantia_hasta: str = Form(""),
    valor_compra: str = Form(""), observaciones: str = Form(""),
    user=Depends(get_current_admin), db: Session = Depends(get_db),
):
    if not user:
        return redirect("/login")
    ficha = db.query(models.ActivoFicha).filter(models.ActivoFicha.equipo_id == equipo_id).first()
    if not ficha:
        ficha = models.ActivoFicha(equipo_id=equipo_id)
        db.add(ficha)
    ficha.placa_inventario = placa_inventario.strip() or None
    ficha.serial = serial.strip() or None
    ficha.estado = estado if estado in ESTADOS_ACTIVO else "Activo"
    ficha.fecha_compra = parse_fecha(fecha_compra)
    ficha.proveedor = proveedor.strip() or None
    ficha.garantia_hasta = parse_fecha(garantia_hasta)
    ficha.valor_compra = parse_int(valor_compra)
    ficha.observaciones = observaciones.strip() or None
    db.commit()
    return redirect(f"/admin/activos/{equipo_id}")


@router.post("/admin/activos/{equipo_id}/mantenimiento")
def activo_agregar_mantenimiento(
    equipo_id: str,
    fecha: str = Form(...), tipo: str = Form(...), descripcion: str = Form(...),
    tecnico: str = Form(""), ticket_id: str = Form(""), costo: str = Form(""),
    user=Depends(get_current_admin), db: Session = Depends(get_db),
):
    if not user:
        return redirect("/login")
    db.add(models.ActivoMantenimiento(
        equipo_id=equipo_id, fecha=parse_fecha(fecha) or hoy(), tipo=tipo, descripcion=descripcion.strip(),
        tecnico=tecnico.strip() or None, ticket_id=parse_int(ticket_id), costo=parse_int(costo),
        registrado_por=user.username, fecha_registro=ahora(),
    ))
    db.commit()
    return redirect(f"/admin/activos/{equipo_id}#mantenimientos")


@router.post("/admin/activos/{equipo_id}/software")
def activo_agregar_software(
    equipo_id: str,
    fecha: str = Form(...), accion: str = Form(...), nombre: str = Form(...),
    version: str = Form(""), licencia: str = Form(""), notas: str = Form(""),
    user=Depends(get_current_admin), db: Session = Depends(get_db),
):
    if not user:
        return redirect("/login")
    db.add(models.ActivoSoftware(
        equipo_id=equipo_id, fecha=parse_fecha(fecha) or hoy(), accion=accion, nombre=nombre.strip(),
        version=version.strip() or None, licencia=licencia.strip() or None, notas=notas.strip() or None,
        registrado_por=user.username, fecha_registro=ahora(),
    ))
    db.commit()
    return redirect(f"/admin/activos/{equipo_id}#software")


@router.post("/admin/activos/registro/{clase}/{registro_id}/delete")
def activo_eliminar_registro(clase: str, registro_id: int, user=Depends(get_current_admin), db: Session = Depends(get_db)):
    if not user:
        return redirect("/login")
    modelo = {"mantenimiento": models.ActivoMantenimiento, "software": models.ActivoSoftware,
              "foto": models.ActivoFoto}.get(clase)
    if not modelo:
        raise HTTPException(status_code=404)
    registro = db.get(modelo, registro_id)
    equipo_id = registro.equipo_id if registro else ""
    if registro:
        db.delete(registro)
        db.commit()
    return redirect(f"/admin/activos/{equipo_id}" if equipo_id else "/admin/activos")


# ---------------------------------------------------------------------------
# Bóveda de contraseñas
# ---------------------------------------------------------------------------

def fernet():
    clave = os.getenv("VAULT_KEY")
    if not clave:
        return None
    try:
        return Fernet(clave.encode())
    except (ValueError, TypeError):
        return None


def cifrar(texto: Optional[str]):
    if not texto:
        return None
    return fernet().encrypt(texto.encode("utf-8")).decode()


def descifrar(texto: Optional[str]):
    if not texto:
        return ""
    try:
        return fernet().decrypt(texto.encode()).decode("utf-8")
    except InvalidToken:
        return "⚠ No se pudo descifrar (la VAULT_KEY cambió)"


def boveda_desbloqueada(request: Request, user):
    if not user:
        return False
    payload = decode_token(request.cookies.get(BOVEDA_COOKIE, ""))
    return bool(payload and payload.get("scope") == "boveda" and payload.get("sub") == user.username)


def registrar_acceso(db: Session, request: Request, user, accion, credencial=None, servicio=None):
    db.add(models.CredencialAcceso(
        credencial_id=credencial.id if credencial else None,
        servicio=servicio or (credencial.servicio if credencial else None),
        usuario_app=user.username if user else "?", accion=accion,
        ip=(request.headers.get("x-forwarded-for") or (request.client.host if request.client else ""))[:100],
        fecha=ahora(),
    ))


ICONOS_CATEGORIA = {
    "Correo": "✉️", "Aplicativo": "🧩", "Equipo": "💻", "Cámara / CCTV": "📹", "Red / Wi-Fi": "📶",
    "Hosting / Dominio": "🌐", "Software": "🗄️", "Otro": "🔑",
}


def _grupo_y_detalle(servicio: str):
    """'Gmail · Producción' -> ('Gmail', 'Producción'); 'Equipo 1742 · Portátil' -> ('Equipo', '1742 · Portátil')."""
    grupo, _, detalle = (servicio or "").partition(" · ")
    base = re.sub(r"\s*\(?[\w-]*\d[\w-]*\)?", "", grupo).strip()
    if base and base != grupo:
        detalle = " · ".join(p for p in (grupo[len(base):].strip(), detalle) if p) if grupo.startswith(base) \
            else " · ".join(p for p in (grupo, detalle) if p)
        grupo = base
    return grupo.strip(), (detalle or grupo).strip()


def _preparar_grupo(nombre, items):
    """Si el detalle se repite (p. ej. 'Comercial' en Gmail), se usa como subtítulo y cada fila muestra el usuario."""
    detalles = [it["detalle"] for it in items]
    con_subtitulos = len(set(detalles)) < len(detalles)
    urls = {(it["c"].url or "") for it in items}
    url_comun = urls.pop() if len(urls) == 1 else ""
    if con_subtitulos:
        items = sorted(items, key=lambda it: (it["detalle"].lower(), (it["c"].usuario or "").lower()))
    anterior = None
    for it in items:
        c = it["c"]
        it["subtitulo"] = it["detalle"] if con_subtitulos and it["detalle"] != anterior else ""
        anterior = it["detalle"]
        it["titulo"] = (c.usuario or it["detalle"]) if con_subtitulos else it["detalle"]
        it["mostrar_usuario"] = bool(c.usuario) and it["titulo"] != c.usuario
        it["mostrar_url"] = bool(c.url) and not url_comun
    return {"nombre": nombre, "lista": items, "url_comun": url_comun}


def _unir_por_primera_palabra(grupos):
    """Une los sueltos con el grupo que comparte su primera palabra ('Avansis Nómina' -> 'Avansis')."""
    def primera(nombre):
        return nombre.split()[0] if nombre.split() else nombre

    def resto(texto, prefijo):
        return texto[len(prefijo):].strip(" ·") if texto.lower().startswith(prefijo.lower()) else texto

    resultado = {}
    sueltos = []
    for nombre, items in grupos.items():
        (resultado.__setitem__(nombre, items) if len(items) > 1 else sueltos.append((nombre, items[0])))

    for nombre, it in sueltos:
        clave = primera(nombre).lower()
        destino = [g for g in resultado if primera(g).lower() == clave]
        companeros = [n for n, _ in sueltos if primera(n).lower() == clave]
        if len(destino) == 1 or (not destino and len(companeros) > 1):
            nuevo = primera(destino[0] if destino else nombre)
            if destino and destino[0] != nuevo:  # renombrar el grupo existente a la palabra común
                anteriores = resultado.pop(destino[0])
                sobrante = resto(destino[0], nuevo)
                for x in anteriores:
                    x["detalle"] = " · ".join(p for p in (sobrante, x["detalle"]) if p)
                resultado[nuevo] = anteriores
            it["detalle"] = resto(it["c"].servicio, nuevo) or it["c"].servicio
            resultado.setdefault(nuevo, []).append(it)
        else:
            resultado.setdefault(nombre, []).append(it)

    # Grupos hermanos pequeños ('Avansis Inventario', 'Avansis Nómina'…) -> un grupo con subtítulos
    por_palabra = {}
    for nombre in resultado:
        por_palabra.setdefault(primera(nombre).lower(), []).append(nombre)
    for nombres in por_palabra.values():
        if len(nombres) < 2 or sum(len(resultado[n]) for n in nombres) > 20:
            continue
        nuevo = primera(nombres[0])
        unidos = []
        for n in nombres:
            for x in resultado.pop(n):
                sub = resto(n, nuevo)
                x["detalle"] = sub if (sub and x["c"].usuario) else " · ".join(p for p in (sub, x["detalle"]) if p)
                unidos.append(x)
        resultado[nuevo] = unidos
    return resultado


def agrupar_credenciales(credenciales):
    """Categorías -> grupos -> credenciales, para mostrar la bóveda ordenada en lugar de una lista plana."""
    por_categoria = {}
    for c in credenciales:
        grupo, detalle = _grupo_y_detalle(c.servicio)
        por_categoria.setdefault(c.categoria or "Otro", {}).setdefault(grupo, []).append({"c": c, "detalle": detalle})

    orden_cat = {nombre: i for i, nombre in enumerate(CATEGORIAS_CREDENCIAL)}
    secciones = []
    for categoria, grupos in por_categoria.items():
        if sum(len(items) for items in grupos.values()) <= 6:  # categorías pequeñas: un solo bloque
            items = [it for lst in grupos.values() for it in lst]
            grupos = {"General": items}
            for it in items:
                it["detalle"] = it["c"].servicio
        grupos = _unir_por_primera_palabra(grupos)
        normales = {g: items for g, items in grupos.items() if len(items) > 1}
        sueltos = [it for g, items in grupos.items() if len(items) == 1 for it in items]
        if sueltos:
            for it in sueltos:  # en el grupo "Otros" se muestra el nombre completo
                it["detalle"] = it["c"].servicio
            normales["Otros" if normales else "General"] = sorted(sueltos, key=lambda it: it["detalle"].lower())
        lista = sorted(normales.items(), key=lambda kv: (kv[0] in ("Otros",), -len(kv[1]), kv[0].lower()))
        secciones.append({
            "nombre": categoria,
            "icono": ICONOS_CATEGORIA.get(categoria, "🔑"),
            "total": sum(len(items) for _, items in lista),
            "grupos": [_preparar_grupo(g, items) for g, items in lista],
        })
    secciones.sort(key=lambda s: (orden_cat.get(s["nombre"], 99), s["nombre"]))
    return secciones


@router.get("/admin/boveda", response_class=HTMLResponse)
def boveda_page(request: Request, error: str = "", user=Depends(get_current_admin), db: Session = Depends(get_db)):
    if not user:
        return redirect("/login")
    contexto = {
        "request": request, "user": user, "error": error, "categorias": CATEGORIAS_CREDENCIAL,
        "configurada": fernet() is not None, "desbloqueada": False,
        "clave_por_defecto": verify_password(DEFAULT_ADMIN_PASSWORD, user.password_hash),
    }
    if contexto["configurada"] and boveda_desbloqueada(request, user):
        credenciales = db.query(models.Credencial).order_by(models.Credencial.categoria, models.Credencial.servicio).all()
        contexto.update({
            "desbloqueada": True,
            "credenciales": credenciales,
            "secciones": agrupar_credenciales(credenciales),
            "accesos": db.query(models.CredencialAcceso).order_by(models.CredencialAcceso.fecha.desc()).limit(60).all(),
            "equipos": {e.get("id"): e.get("nombre") for e in cargar_equipos(db)},
            "minutos": BOVEDA_MINUTOS,
            "segundos_restantes": max(0, int(decode_token(request.cookies.get(BOVEDA_COOKIE, "")).get("exp", 0)
                                             - datetime.now(timezone.utc).timestamp())),
        })
    return templates.TemplateResponse(request=request, name="ti_boveda.html", context=contexto)


@router.post("/admin/boveda/desbloquear")
def boveda_desbloquear(request: Request, password: str = Form(...), user=Depends(get_current_admin), db: Session = Depends(get_db)):
    if not user:
        return redirect("/login")
    if verify_password(DEFAULT_ADMIN_PASSWORD, user.password_hash):
        return redirect("/admin/boveda?error=default")
    if not verify_password(password, user.password_hash):
        registrar_acceso(db, request, user, "Intento fallido")
        db.commit()
        return redirect("/admin/boveda?error=clave")
    registrar_acceso(db, request, user, "Desbloquear")
    db.commit()
    response = redirect("/admin/boveda")
    token = create_access_token({"sub": user.username, "scope": "boveda"}, minutes=BOVEDA_MINUTOS)
    response.set_cookie(BOVEDA_COOKIE, token, httponly=True, samesite="strict", secure=COOKIE_SECURE,
                        path=BOVEDA_PATH, max_age=BOVEDA_MINUTOS * 60)
    return response


@router.post("/admin/boveda/bloquear")
def boveda_bloquear():
    response = redirect("/admin/boveda")
    response.delete_cookie(BOVEDA_COOKIE, path=BOVEDA_PATH)
    return response


@router.post("/admin/boveda/{credencial_id}/revelar")
def boveda_revelar(request: Request, credencial_id: int, user=Depends(get_current_admin), db: Session = Depends(get_db)):
    if not user or not fernet() or not boveda_desbloqueada(request, user):
        return JSONResponse({"error": "Bóveda bloqueada"}, status_code=401)
    credencial = db.get(models.Credencial, credencial_id)
    if not credencial:
        return JSONResponse({"error": "No existe"}, status_code=404)
    registrar_acceso(db, request, user, "Ver", credencial)
    db.commit()
    return {"clave": descifrar(credencial.clave_cifrada), "notas": descifrar(credencial.notas_cifradas)}


@router.post("/admin/boveda/guardar")
def boveda_guardar(
    request: Request,
    credencial_id: Optional[int] = Form(None),
    categoria: str = Form(...), servicio: str = Form(...), usuario: str = Form(""),
    clave: str = Form(""), url: str = Form(""), equipo_id: str = Form(""), notas: str = Form(""),
    user=Depends(get_current_admin), db: Session = Depends(get_db),
):
    if not user:
        return redirect("/login")
    if not fernet() or not boveda_desbloqueada(request, user):
        return redirect("/admin/boveda")
    credencial = db.get(models.Credencial, credencial_id) if credencial_id else None
    nueva = credencial is None
    if nueva:
        credencial = models.Credencial(fecha_creacion=ahora())
        db.add(credencial)
    credencial.categoria = categoria
    credencial.servicio = servicio.strip()
    credencial.usuario = usuario.strip() or None
    credencial.url = url.strip() or None
    credencial.equipo_id = equipo_id or None
    if clave or nueva:
        credencial.clave_cifrada = cifrar(clave)
    if notas or nueva:
        credencial.notas_cifradas = cifrar(notas.strip())
    credencial.actualizado_por = user.username
    credencial.fecha_actualizacion = ahora()
    db.flush()
    registrar_acceso(db, request, user, "Crear" if nueva else "Editar", credencial)
    db.commit()
    return redirect("/admin/boveda")


@router.post("/admin/boveda/{credencial_id}/delete")
def boveda_eliminar(request: Request, credencial_id: int, user=Depends(get_current_admin), db: Session = Depends(get_db)):
    if not user:
        return redirect("/login")
    if not boveda_desbloqueada(request, user):
        return redirect("/admin/boveda")
    credencial = db.get(models.Credencial, credencial_id)
    if credencial:
        registrar_acceso(db, request, user, "Eliminar", credencial)
        db.delete(credencial)
        db.commit()
    return redirect("/admin/boveda")


@router.post("/admin/boveda/importar")
async def boveda_importar(request: Request, archivo: UploadFile = File(...), user=Depends(get_current_admin), db: Session = Depends(get_db)):
    if not user:
        return redirect("/login")
    if not fernet() or not boveda_desbloqueada(request, user):
        return redirect("/admin/boveda")
    contenido = (await archivo.read()).decode("utf-8-sig", errors="replace")
    try:
        dialecto = csv.Sniffer().sniff(contenido.splitlines()[0] if contenido else ",", delimiters=",;\t")
    except csv.Error:
        dialecto = csv.excel
    lector = csv.DictReader(io.StringIO(contenido), dialect=dialecto)
    total = 0
    for fila in lector:
        fila = {(k or "").strip().lower(): (v or "").strip() for k, v in fila.items()}
        servicio = fila.get("servicio") or fila.get("nombre")
        if not servicio:
            continue
        categoria = fila.get("categoria") or fila.get("categoría") or "Otro"
        db.add(models.Credencial(
            categoria=categoria, servicio=servicio,
            usuario=fila.get("usuario") or None,
            clave_cifrada=cifrar(fila.get("clave") or fila.get("contraseña") or fila.get("password")),
            url=fila.get("url") or None,
            notas_cifradas=cifrar(fila.get("notas")),
            actualizado_por=user.username, fecha_creacion=ahora(), fecha_actualizacion=ahora(),
        ))
        total += 1
    registrar_acceso(db, request, user, "Importar", servicio=f"{total} credenciales desde {archivo.filename}")
    db.commit()
    return redirect("/admin/boveda")


@router.get("/admin/cuenta", response_class=HTMLResponse)
def cuenta_page(request: Request, msg: str = "", user=Depends(get_current_admin)):
    if not user:
        return redirect("/login")
    return templates.TemplateResponse(request=request, name="ti_cuenta.html", context={
        "request": request, "user": user, "msg": msg,
        "clave_por_defecto": verify_password(DEFAULT_ADMIN_PASSWORD, user.password_hash),
    })


@router.post("/admin/cuenta")
def cuenta_cambiar_clave(
    actual: str = Form(...), nueva: str = Form(...), confirmar: str = Form(...),
    user=Depends(get_current_admin), db: Session = Depends(get_db),
):
    if not user:
        return redirect("/login")
    if not verify_password(actual, user.password_hash):
        return redirect("/admin/cuenta?msg=actual")
    if nueva != confirmar:
        return redirect("/admin/cuenta?msg=distintas")
    if len(nueva) < 10 or nueva == DEFAULT_ADMIN_PASSWORD:
        return redirect("/admin/cuenta?msg=corta")
    user.password_hash = get_password_hash(nueva)
    db.commit()
    return redirect("/admin/cuenta?msg=ok")


# ---------------------------------------------------------------------------
# Dominios, hosting y vencimientos
# ---------------------------------------------------------------------------

def nivel_dominio(dias):
    if dias is None:
        return "sin_fecha"
    if dias < 0:
        return "vencido"
    if dias <= 7:
        return "critico"
    if dias <= 30:
        return "alto"
    if dias <= 60:
        return "medio"
    return "ok"


@router.get("/admin/dominios", response_class=HTMLResponse)
def dominios_page(request: Request, user=Depends(get_current_admin), db: Session = Depends(get_db)):
    if not user:
        return redirect("/login")
    dominios = db.query(models.Dominio).order_by(models.Dominio.fecha_vencimiento).all()
    filas = []
    for d in dominios:
        dias = (d.fecha_vencimiento - hoy()).days if d.fecha_vencimiento else None
        filas.append({"d": d, "dias": dias, "nivel": nivel_dominio(dias)})
    return templates.TemplateResponse(request=request, name="ti_dominios.html", context={
        "request": request, "user": user, "filas": filas, "hoy": hoy(),
    })


@router.post("/admin/dominios/guardar")
def dominios_guardar(
    dominio_id: Optional[int] = Form(None),
    nombre: str = Form(...), tipo: str = Form("Dominio"), proveedor: str = Form(""),
    fecha_vencimiento: str = Form(...), renovacion_automatica: Optional[str] = Form(None),
    responsable: str = Form(""), costo_anual: str = Form(""), servicios: str = Form(""), notas: str = Form(""),
    user=Depends(get_current_admin), db: Session = Depends(get_db),
):
    if not user:
        return redirect("/login")
    dominio = db.get(models.Dominio, dominio_id) if dominio_id else None
    if not dominio:
        dominio = models.Dominio()
        db.add(dominio)
    dominio.nombre = nombre.strip().lower()
    dominio.tipo = tipo
    dominio.proveedor = proveedor.strip() or None
    dominio.fecha_vencimiento = parse_fecha(fecha_vencimiento)
    dominio.renovacion_automatica = bool(renovacion_automatica)
    dominio.responsable = responsable.strip() or None
    dominio.costo_anual = parse_int(costo_anual)
    dominio.servicios = servicios.strip() or None
    dominio.notas = notas.strip() or None
    db.commit()
    return redirect("/admin/dominios")


@router.post("/admin/dominios/{dominio_id}/renovar")
def dominios_renovar(dominio_id: int, anios: int = Form(1), user=Depends(get_current_admin), db: Session = Depends(get_db)):
    if not user:
        return redirect("/login")
    dominio = db.get(models.Dominio, dominio_id)
    if dominio and dominio.fecha_vencimiento:
        f = dominio.fecha_vencimiento
        try:
            dominio.fecha_vencimiento = f.replace(year=f.year + max(1, anios))
        except ValueError:  # 29 de febrero
            dominio.fecha_vencimiento = f.replace(year=f.year + max(1, anios), day=28)
        db.commit()
    return redirect("/admin/dominios")


@router.post("/admin/dominios/{dominio_id}/delete")
def dominios_eliminar(dominio_id: int, user=Depends(get_current_admin), db: Session = Depends(get_db)):
    if not user:
        return redirect("/login")
    dominio = db.get(models.Dominio, dominio_id)
    if dominio:
        db.delete(dominio)
        db.commit()
    return redirect("/admin/dominios")


# ---------------------------------------------------------------------------
# Alertas para el panel principal
# ---------------------------------------------------------------------------

def calcular_alertas(db: Session):
    alertas = []
    for d in db.query(models.Dominio).all():
        if not d.fecha_vencimiento:
            continue
        dias = (d.fecha_vencimiento - hoy()).days
        nivel = nivel_dominio(dias)
        if nivel == "ok":
            continue
        if dias < 0:
            texto = f"venció hace {-dias} día(s)"
        elif dias == 0:
            texto = "vence HOY"
        else:
            texto = f"vence en {dias} día(s) ({d.fecha_vencimiento.strftime('%d/%m/%Y')})"
        auto = " · renovación automática activa" if d.renovacion_automatica else ""
        alertas.append({
            "nivel": "critico" if nivel in ("vencido", "critico") else nivel,
            "titulo": f"{d.tipo} {d.nombre} {texto}",
            "detalle": f"Proveedor: {d.proveedor or 'sin registrar'}{auto}",
            "url": "/admin/dominios",
        })

    problemas = []
    for e in db.query(models.RespaldoEstacion).filter(models.RespaldoEstacion.activo == True).all():  # noqa: E712
        r = resumen_estacion(e)
        if r["nivel"] in ("vencido", "error", "sin_datos"):
            problemas.append(e.nombre or e.hostname)
    if problemas:
        listado = ", ".join(problemas[:5]) + (f" y {len(problemas) - 5} más" if len(problemas) > 5 else "")
        alertas.append({
            "nivel": "alto",
            "titulo": f"{len(problemas)} estación(es) sin respaldo válido",
            "detalle": listado,
            "url": "/admin/respaldos",
        })

    orden = {"critico": 0, "alto": 1, "medio": 2}
    alertas.sort(key=lambda a: orden.get(a["nivel"], 3))
    return alertas
